from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from ulsan_jobs.gangsaitda import HOLD, SCHEDULE_FALLBACK, build_gangsaitda, clean_title, region, row_values
from ulsan_jobs.models import Posting

TEMPLATE = Path(__file__).resolve().parent.parent / "config" / "gangsaitda_template.xlsx"
HEADERS = ["제목", "기관명", "지역", "마감일", "수업 일정", "상세 내용", "수업 대상", "모집 인원",
           "지원 자격", "제출 서류", "원문 링크", "지원서 링크", "지원 이메일"]


def test_clean_title_drops_dates_and_periods():
    assert clean_title("2026학년도 메아리학교 초등 방과후 맞춤형 프로그램(도담이음) 개인 위탁 외부 강사 공모") == (
        "메아리학교 초등 방과후 맞춤형 프로그램(도담이음) 개인 위탁 외부 강사 공모"
    )
    assert clean_title("[울주종합체육센터]2026년 하반기 교육(체육)강사 긴급 위·수탁 모집 공고") == (
        "[울주종합체육센터] 하반기 교육(체육)강사 긴급 위·수탁 모집 공고"
    )
    assert clean_title("청소년센터 강사 모집 공고(~9.30까지)") == "청소년센터 강사 모집 공고"
    assert clean_title("동천고등학교 시간강사 채용 공고(일본어) 10.23-11.20") == "동천고등학교 시간강사 채용 공고(일본어)"
    assert clean_title("2026년 울산광역시여성회관 교육강사(3학기 단기특강) 모집 공고") == (
        "울산광역시여성회관 교육강사(3학기 단기특강) 모집 공고"
    )


def test_region_uses_district_or_org_name():
    assert region(Posting("s", "t", "u", district="남구")) == "울산 남구"
    assert region(Posting("s", "강사 모집", "u", org_name="울산광역시남구도시관리공단", district="울산전체")) == "울산 남구"
    assert region(Posting("s", "강사 모집", "u", org_name="척과초등학교", district="울산전체")) == "울산"


def test_row_fills_required_cells_and_holds_without_deadline():
    found = Posting("use", "2026학년도 방과후 로봇 강사 모집", "https://example.org/1", org_name="척과초등학교",
                    deadline=date(2026, 10, 1), category="방과후·늘봄",
                    info={"schedule": "2026.10.6.~12.15.", "headcount": 1, "email": "job@example.org"})
    row = row_values(found, "울산교육청 통합인력풀 - 일반채용공고(개인위탁 포함)")
    assert row["제목"] == "방과후 로봇 강사 모집"
    assert row["마감일"] == "2026-10-01" and row["처리"] == ""
    assert row["수업 일정"] == "2026.10.6.~12.15." and row["모집 인원"] == 1 and row["지원 이메일"] == "job@example.org"
    assert "분야: 방과후·늘봄" in row["상세 내용"] and "원문 공고: https://example.org/1" in row["상세 내용"]
    assert row["메모"] == "출처: 울산교육청 통합인력풀 - 일반채용공고(개인위탁 포함)"

    unknown = Posting("uic", "수영 강사 모집", "https://example.org/2", org_name="")
    row = row_values(unknown, "울산시설공단 - 강습위탁")
    assert row["처리"] == HOLD and row["마감일"] == ""
    assert row["수업 일정"] == SCHEDULE_FALLBACK
    assert row["기관명"] == "울산시설공단"  # 기관명이 없으면 게시판 이름에서
    assert row["메모"].startswith("마감일을 찾지 못해 보류")


def test_workbook_keeps_template_and_adds_rows(tmp_path):
    postings = [
        Posting("use", "방과후 로봇 강사 모집", "https://example.org/1", org_name="척과초등학교", deadline=date(2026, 10, 1)),
        Posting("uic", "수영 강사 모집", "https://example.org/2", org_name="울산시설공단"),
    ]
    path = build_gangsaitda(tmp_path / "강사잇다.xlsx", postings, {"use": "교육청", "uic": "시설공단"}, TEMPLATE)
    wb = load_workbook(path)
    assert wb.sheetnames == ["공고", "예시", "안내"] and wb.active.title == "공고"
    ws = wb["공고"]
    headers = [c.value for c in ws[1]]
    assert headers == HEADERS + ["처리", "메모"]
    assert ws["A1"].comment is not None  # 양식의 칸 설명은 그대로
    assert ws["A2"].value == "방과후 로봇 강사 모집" and ws["D2"].value == "2026-10-01"
    assert ws["K2"].hyperlink.target == "https://example.org/1"
    assert ws["N2"].value is None and ws["N3"].value == HOLD
    assert ws["F2"].alignment.wrap_text
