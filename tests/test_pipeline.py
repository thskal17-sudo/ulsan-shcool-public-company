import shutil
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ulsan_jobs import pipeline
from ulsan_jobs.collectors import COLLECTORS, Collector
from ulsan_jobs.models import KST, Posting
from ulsan_jobs.storage import Store

ROOT = Path(__file__).resolve().parent.parent

SOURCES_YAML = """
max_phase: 1
sources:
  - id: fake
    name: 가짜 게시판
    org_type: 교육청·학교
    district: 울산전체
    collector: fake
    url: https://example.org/list
    keyword_filter: true
    phase: 1
    enabled: true
  - id: broken
    name: 고장난 게시판
    org_type: 지자체
    collector: fake_broken
    url: https://example.org/broken
    phase: 1
    enabled: true
  - id: later
    name: 다음 단계 게시판
    collector: fake
    url: https://example.org/later
    phase: 2
    enabled: true
"""

POSTS = [
    # 신규·마감 전
    Posting("fake", "방과후 독서논술 강사 모집", "https://example.org/1", "1", "척과초등학교", deadline=date(2026, 10, 1)),
    # 마감임박
    Posting("fake", "수영 강사 모집", "https://example.org/2", "2", "울산시설공단", deadline=date(2026, 9, 28)),
    # 이미 마감 → 신규 아님
    Posting("fake", "늘봄 강사 모집", "https://example.org/3", "3", "남목초등학교", deadline=date(2026, 9, 20)),
    # 강사 공고 아님
    Posting("fake", "기간제 교원 채용", "https://example.org/4", "4", "대현중학교"),
]


class FakeCollector(Collector):
    def collect(self):
        return [Posting(**{**p.__dict__}) for p in POSTS]


class BrokenCollector(Collector):
    def collect(self):
        raise ConnectionError("접속 실패")


@pytest.fixture
def env(tmp_path, monkeypatch):
    config = tmp_path / "config"
    config.mkdir()
    (config / "sources.yaml").write_text(SOURCES_YAML, encoding="utf-8")
    shutil.copy(ROOT / "config" / "keywords.yaml", config / "keywords.yaml")
    monkeypatch.setitem(COLLECTORS, "fake", FakeCollector)
    monkeypatch.setitem(COLLECTORS, "fake_broken", BrokenCollector)
    sent = []
    monkeypatch.setattr(pipeline, "send", lambda cfg, msg: sent.append(msg))
    monkeypatch.setenv("SMTP_USER", "me@example.org")
    monkeypatch.setenv("SMTP_APP_PASSWORD", "abcd efgh")
    return tmp_path, config, sent


NOW = datetime(2026, 9, 26, 7, 0, tzinfo=KST)


def run(tmp_path, config, send_mail, now=NOW):
    return pipeline.run(
        db_path=tmp_path / "db.sqlite", out_dir=tmp_path / "out", send_mail=send_mail, config_dir=config, now=now
    )


def test_run_builds_report_and_mails(env):
    tmp_path, config, sent = env
    outcome = run(tmp_path, config, send_mail=True)

    states = {r.source_id: r.state for r in outcome.results}
    assert states == {"fake": "정상", "broken": "오류"}  # phase 2 소스는 수집하지 않음
    assert [p.title for p in outcome.new] == ["수영 강사 모집", "방과후 독서논술 강사 모집"]  # 마감 빠른 순
    assert [p.title for p in outcome.closing_soon] == ["수영 강사 모집"]
    assert outcome.mailed and len(sent) == 1

    msg = sent[0]
    assert "신규 2건" in msg["Subject"] and "마감임박 1건" in msg["Subject"]
    # 메일에는 강사잇다 올리기 양식만 붙는다 (마감 전 공고 전부)
    [attachment] = list(msg.iter_attachments())
    assert attachment.get_filename() == "강사잇다_울산_2026-09-26.xlsx"
    assert "강사잇다 올리기 양식입니다 (마감 전 공고 2건)" in msg.get_body(("html",)).get_content()
    upload = load_workbook(outcome.upload_path)["공고"]
    assert [upload.cell(row=r, column=1).value for r in (2, 3)] == ["수영 강사 모집", "방과후 독서논술 강사 모집"]
    assert upload["D2"].value == "2026-09-28"

    wb = load_workbook(outcome.report_path)
    assert wb.sheetnames == ["신규 (2)", "마감임박 (1)", "진행중 전체 (2)", "수집현황"]
    ws = wb["신규 (2)"]
    assert ws["F2"].value == "수영 강사 모집"
    assert ws["F2"].hyperlink.target == "https://example.org/2"
    assert ws["I2"].value == "D-2"
    status = wb["수집현황"]
    assert status["C3"].value == "오류"


def test_second_run_has_no_new_after_mail(env):
    tmp_path, config, sent = env
    run(tmp_path, config, send_mail=True)
    outcome = run(tmp_path, config, send_mail=True, now=datetime(2026, 9, 27, 7, 0, tzinfo=KST))
    assert outcome.new == []
    assert len(sent) == 2  # 신규가 없어도 확인용 메일은 보냄


def test_unsent_postings_stay_new(env):
    # 메일을 보내지 않은(또는 실패한) 실행 뒤에는 다음 실행에서 다시 신규로 잡혀야 한다
    tmp_path, config, sent = env
    run(tmp_path, config, send_mail=False)
    outcome = run(tmp_path, config, send_mail=True)
    assert len(outcome.new) == 2


def test_missing_mail_settings_fail_loudly(env, monkeypatch):
    tmp_path, config, _ = env
    monkeypatch.delenv("SMTP_APP_PASSWORD")
    with pytest.raises(RuntimeError, match="SMTP_APP_PASSWORD"):
        run(tmp_path, config, send_mail=True)


class FlakyCollector(Collector):
    calls = 0

    def collect(self):
        FlakyCollector.calls += 1
        if FlakyCollector.calls == 1:
            raise ConnectionError("일시 장애")
        return [Posting(**{**p.__dict__}) for p in POSTS]


def test_failed_source_is_retried_once(env, monkeypatch):
    # 관공서 서버의 일시 장애: 처음 실패한 소스는 끝에서 한 번 더 시도한다
    tmp_path, config, _ = env
    FlakyCollector.calls = 0
    monkeypatch.setitem(COLLECTORS, "fake", FlakyCollector)
    outcome = run(tmp_path, config, send_mail=False)
    states = {r.source_id: r.state for r in outcome.results}
    assert states == {"fake": "정상", "broken": "오류"}
    assert FlakyCollector.calls == 2
    assert len(outcome.new) == 2


class _Clock:
    def __init__(self):
        self.now = 0.0

    def __call__(self):
        return self.now


class _RecordingHttp:
    def __init__(self):
        self.forgets = 0

    def forget_unreachable(self):
        self.forgets += 1


def _collect(tmp_path, config, clock, http):
    from ulsan_jobs.config import load_rules, load_settings
    from ulsan_jobs.storage import Store

    sources = pipeline.select_sources(load_settings(config).sources, None, 1)  # fake, broken
    store = Store(tmp_path / "db.sqlite")
    try:
        return pipeline.collect_all(sources, load_rules(config), store, http, NOW, time_limit=20 * 60, clock=clock)
    finally:
        store.close()


def test_retry_pass_forgets_unreachable_hosts(env):
    # 끝의 재시도에서는 앞에서 접속 실패한 서버도 한 번 더 기다려 본다
    tmp_path, config, _ = env
    http = _RecordingHttp()
    results = _collect(tmp_path, config, _Clock(), http)
    assert [r.state for r in results] == ["정상", "오류"]
    assert http.forgets == 1


def test_time_limit_skips_remaining_sources_and_retries(env, monkeypatch):
    # 접속 안 되는 서버가 많아 수집이 오래 걸리면 남은 소스·재시도를 건너뛰고 메일을 보낼 수 있게 한다
    tmp_path, config, _ = env
    clock = _Clock()
    broken_calls = []

    class SlowCollector(FakeCollector):
        def collect(self):
            clock.now += 25 * 60
            return super().collect()

    class CountingBroken(BrokenCollector):
        def collect(self):
            broken_calls.append(1)
            return super().collect()

    monkeypatch.setitem(COLLECTORS, "fake", SlowCollector)
    monkeypatch.setitem(COLLECTORS, "fake_broken", CountingBroken)
    http = _RecordingHttp()
    fake, broken = _collect(tmp_path, config, clock, http)

    assert fake.state == "정상"
    assert broken.state == "시간초과" and "20분을 넘겨 건너뜀" in broken.error
    assert broken.needs_other_server  # 보충 수집 대상
    assert broken_calls == [] and http.forgets == 0


# ─────────────────────────────── 실패 설명 · 보충 수집


def test_failure_descriptions_are_plain():
    import requests
    from urllib3.exceptions import MaxRetryError, NewConnectionError

    from ulsan_jobs.http import HostUnreachable

    def http_error(code):
        resp = requests.Response()
        resp.status_code = code
        return requests.exceptions.HTTPError(response=resp)

    refused = requests.exceptions.ConnectionError(MaxRetryError(None, "/", NewConnectionError(None, "refused")))
    cases = {
        requests.exceptions.ConnectTimeout("timed out"): "접속불가",
        refused: "접속불가",
        HostUnreachable("x"): "접속불가",
        http_error(500): "오류",
        requests.exceptions.ReadTimeout("slow"): "오류",
        ValueError("표 모양이 다름"): "오류",
    }
    for exc, state in cases.items():
        got_state, message = pipeline.describe_failure(exc)
        assert got_state == state, exc
        assert "HTTPSConnectionPool" not in message and "Traceback" not in message
    assert "HTTP 500" in pipeline.describe_failure(http_error(500))[1]
    assert "해외 접속 차단" in pipeline.describe_failure(requests.exceptions.ConnectTimeout("t"))[1]


class UnreachableCollector(Collector):
    calls = 0

    def collect(self):
        import requests

        UnreachableCollector.calls += 1
        raise requests.exceptions.ConnectTimeout("timed out")


def test_unreachable_source_is_retried_and_listed_for_other_server(env, monkeypatch, tmp_path):
    # 접속 안 되는 소스는 끝에서 한 번 더 시도하고, 그래도 안 되면 다른 서버에서 다시 수집할 목록에 오른다
    from ulsan_jobs.__main__ import main

    _, config, _ = env
    UnreachableCollector.calls = 0
    monkeypatch.setitem(COLLECTORS, "fake_broken", UnreachableCollector)
    retry_file = tmp_path / "retry.txt"
    code = main([
        "--config", str(config), "run", "--no-mail",
        "--db", str(tmp_path / "db.sqlite"), "--out", str(tmp_path / "out"), "--retry-list", str(retry_file),
    ])
    assert code == 0
    assert UnreachableCollector.calls == 2
    assert retry_file.read_text(encoding="utf-8") == "broken"


def test_catch_up_mails_only_when_something_new(env):
    tmp_path, config, sent = env
    # 아침 실행: 'fake' 공고 2건을 메일로 보냄
    run(tmp_path, config, send_mail=True)
    assert len(sent) == 1

    # 보충 실행: 같은 소스를 다시 읽어도 새 공고가 없으면 메일을 보내지 않는다
    outcome = pipeline.run(
        db_path=tmp_path / "db.sqlite", out_dir=tmp_path / "out", send_mail=True, config_dir=config,
        now=NOW, only=["fake"], catch_up=True,
    )
    assert outcome.new == [] and not outcome.mailed and len(sent) == 1


def test_catch_up_mail_has_only_the_new_postings(env):
    tmp_path, config, sent = env
    # 아침에 'fake' 가 접속 안 됐다고 치고, 보충 실행에서 처음 읽음
    outcome = pipeline.run(
        db_path=tmp_path / "db.sqlite", out_dir=tmp_path / "out", send_mail=True, config_dir=config,
        now=NOW, only=["fake"], catch_up=True,
    )
    assert outcome.mailed and len(sent) == 1
    msg = sent[0]
    assert "보충 신규 2건" in msg["Subject"] and "마감임박" not in msg["Subject"]
    assert next(msg.iter_attachments()).get_filename() == "강사잇다_울산_2026-09-26_보충.xlsx"
    html = msg.get_body(("html",)).get_content()
    assert "다른 수집 서버에서 다시 읽어" in html and "<h3>마감임박" not in html


def test_mail_separates_unreachable_sites_from_real_problems():
    from ulsan_jobs.mailer import html_body
    from ulsan_jobs.models import SourceResult

    results = [
        SourceResult("gojobs", "나라일터 모집공고 (기관명에 '울산' 포함)", state="접속불가", error="서버가 응답하지 않음"),
        SourceResult("junggu", "중구청 채용공고(새올)", state="접속불가", error="서버가 응답하지 않음"),
        SourceResult("uic", "울산시설공단 - 강습위탁", state="0건", error="목록에서 글을 하나도 읽지 못함"),
        SourceResult("ok", "울산도서관 - 공지사항"),
    ]
    html = html_body(date(2026, 9, 26), [], [], 0, results)
    assert "확인이 필요한 사이트 1곳" in html and "울산시설공단 - 강습위탁: 목록에서 글을 하나도 읽지 못함" in html
    assert "접속 안 된 사이트 2곳" in html and "나라일터 모집공고 · 중구청 채용공고(새올)" in html
    assert "기관명에" not in html and "울산도서관" not in html


def run_scheduled(tmp_path, config, now, send_mail=True):
    return pipeline.run(
        db_path=tmp_path / "db.sqlite", out_dir=tmp_path / "out", send_mail=send_mail, config_dir=config,
        now=now, once_daily=True,
    )


class MorePostsCollector(Collector):
    def collect(self):
        extra = Posting("fake", "늘봄 코딩 강사 모집", "https://example.org/5", "5", "성안초등학교",
                        deadline=date(2026, 10, 5))
        return [Posting(**{**p.__dict__}) for p in POSTS] + [extra]


def test_backup_run_sends_nothing_when_nothing_new(env):
    # 새벽 실행이 정기 메일을 보냈으면, 예비 실행은 새 공고가 없을 때 메일을 보내지 않는다
    tmp_path, config, sent = env
    first = run_scheduled(tmp_path, config, datetime(2026, 9, 26, 5, 10, tzinfo=KST))
    assert first.mailed and not first.already_mailed and len(sent) == 1
    assert "마감임박 1건" in sent[0]["Subject"]

    backup = run_scheduled(tmp_path, config, datetime(2026, 9, 26, 8, 20, tzinfo=KST))
    assert backup.already_mailed and backup.new == [] and not backup.mailed and len(sent) == 1


def test_backup_run_mails_only_new_postings_as_supplement(env, monkeypatch):
    tmp_path, config, sent = env
    run_scheduled(tmp_path, config, datetime(2026, 9, 26, 5, 10, tzinfo=KST))
    monkeypatch.setitem(COLLECTORS, "fake", MorePostsCollector)

    backup = run_scheduled(tmp_path, config, datetime(2026, 9, 26, 8, 20, tzinfo=KST))
    assert backup.mailed and [p.title for p in backup.new] == ["늘봄 코딩 강사 모집"] and len(sent) == 2
    msg = sent[1]
    assert "보충 신규 1건" in msg["Subject"] and "마감임박" not in msg["Subject"]
    assert next(msg.iter_attachments()).get_filename() == "강사잇다_울산_2026-09-26_보충.xlsx"
    html = msg.get_body(("html",)).get_content()
    assert "예비 수집에서 새로 찾은 공고 <b>1</b>건" in html and "<h3>마감임박" not in html


def test_backup_run_sends_the_daily_mail_when_the_first_run_did_not(env):
    # 새벽 실행이 통째로 빠졌거나 메일을 못 보냈으면 예비 실행이 정기 메일을 보낸다
    tmp_path, config, sent = env
    run_scheduled(tmp_path, config, datetime(2026, 9, 26, 5, 10, tzinfo=KST), send_mail=False)
    backup = run_scheduled(tmp_path, config, datetime(2026, 9, 26, 8, 20, tzinfo=KST))
    assert backup.mailed and not backup.already_mailed and len(sent) == 1
    assert "신규 2건 · 마감임박 1건" in sent[0]["Subject"]


def test_daily_mail_is_sent_again_the_next_day(env):
    tmp_path, config, sent = env
    run_scheduled(tmp_path, config, datetime(2026, 9, 26, 5, 10, tzinfo=KST))
    next_day = run_scheduled(tmp_path, config, datetime(2026, 9, 27, 5, 10, tzinfo=KST))
    assert next_day.mailed and not next_day.already_mailed and len(sent) == 2
    assert "보충" not in sent[1]["Subject"]


def test_manual_run_always_sends_the_full_mail(env):
    # 손으로 돌린 실행(once_daily 없음)은 오늘 이미 보냈어도 전체 메일을 보낸다
    tmp_path, config, sent = env
    run_scheduled(tmp_path, config, datetime(2026, 9, 26, 5, 10, tzinfo=KST))
    run(tmp_path, config, send_mail=True, now=datetime(2026, 9, 26, 14, 0, tzinfo=KST))
    assert len(sent) == 2 and "보충" not in sent[1]["Subject"]


DECIDED = Posting(
    "fake", "[남부청소년수련관] 주말 수영 강사(긴급) 위·수탁 대상 결정 공고", "https://example.org/9", "9", "울주군시설관리공단"
)


class DecidedCollector(Collector):
    def collect(self):
        return [Posting(**{**p.__dict__}) for p in POSTS] + [Posting(**{**DECIDED.__dict__})]


def test_posting_stored_before_a_new_result_word_is_closed(env, monkeypatch):
    # '대상 결정 공고'를 결과공고로 보기 전에 모집공고로 저장된 글은 다음 수집 때 닫혀 진행중에서 빠진다
    tmp_path, config, sent = env
    store = Store(tmp_path / "db.sqlite")
    store.upsert(Posting(**{**DECIDED.__dict__}), NOW)
    store.close()
    monkeypatch.setitem(COLLECTORS, "fake", DecidedCollector)

    outcome = run(tmp_path, config, send_mail=False)
    assert DECIDED.title not in [p.title for p in outcome.active + outcome.new]
    assert [p.title for p in outcome.new] == ["수영 강사 모집", "방과후 독서논술 강사 모집"]


VOLUNTEER = Posting("fake", "2026학년도 늘봄학교 자원봉사자 모집 공고", "https://example.org/10", "10", "구영초등학교")


class VolunteerCollector(Collector):
    def collect(self):
        return [Posting(**{**p.__dict__}) for p in POSTS] + [Posting(**{**VOLUNTEER.__dict__})]


def test_posting_stored_before_a_new_exclude_word_is_dropped(env, monkeypatch):
    # '봉사자'를 제외 낱말로 넣기 전에 모집공고로 저장된 글은 다음 수집 때 진행중에서 빠진다
    tmp_path, config, sent = env
    store = Store(tmp_path / "db.sqlite")
    store.upsert(Posting(**{**VOLUNTEER.__dict__}), NOW)
    store.close()
    monkeypatch.setitem(COLLECTORS, "fake", VolunteerCollector)

    outcome = run(tmp_path, config, send_mail=False)
    assert VOLUNTEER.title not in [p.title for p in outcome.active + outcome.new]
    store = Store(tmp_path / "db.sqlite")
    assert store.detail_state(VOLUNTEER.uid)[2] == "제외"
    store.close()
