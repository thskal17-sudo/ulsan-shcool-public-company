"""엑셀 리포트: 신규 / 마감임박 / 진행중 전체 / 수집현황 4개 시트."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .dates import days_left, dday_label
from .models import Posting, SourceResult

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
URGENT_FILL = PatternFill("solid", fgColor="F8CBAD")  # D-1 이하
SOON_FILL = PatternFill("solid", fgColor="FFE699")  # D-3 이하
PROBLEM_FILL = PatternFill("solid", fgColor="F8CBAD")
UNREACHABLE_FILL = PatternFill("solid", fgColor="E7E6E6")  # 접속 안 됨 (다음 실행 때 다시 수집)
LINK_FONT = Font(color="0563C1", underline="single")

POSTING_COLUMNS = [
    ("No", 6), ("분야", 13), ("기관유형", 15), ("기관명", 24), ("지역", 8),
    ("제목", 70), ("게시일", 12), ("마감일", 12), ("D-day", 9), ("출처", 30),
]
STATUS_COLUMNS = [
    ("소스", 36), ("기관유형", 15), ("상태", 10), ("목록 글 수", 11), ("강사 공고", 10),
    ("신규", 8), ("비고", 60), ("주소", 60),
]


def sort_key(p: Posting, today: date):
    left = days_left(p.deadline, today)
    return (left is None, left if left is not None else 0, -(p.posted_date or date.min).toordinal())


def build_report(
    path: Path,
    today: date,
    new: list[Posting],
    closing_soon: list[Posting],
    active: list[Posting],
    results: list[SourceResult],
    source_names: dict[str, str],
) -> Path:
    wb = Workbook()
    sheets = [
        (f"신규 ({len(new)})", new),
        (f"마감임박 ({len(closing_soon)})", closing_soon),
        (f"진행중 전체 ({len(active)})", active),
    ]
    for i, (title, postings) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        _write_postings(ws, sorted(postings, key=lambda p: sort_key(p, today)), today, source_names)
    _write_status(wb.create_sheet("수집현황"), results)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _header(ws, columns) -> None:
    for col, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _write_postings(ws, postings: list[Posting], today: date, source_names: dict[str, str]) -> None:
    _header(ws, POSTING_COLUMNS)
    for n, p in enumerate(postings, start=1):
        row = n + 1
        values = [
            n, p.category, p.org_type, p.org_name, p.district, p.title,
            p.posted_date, p.deadline, dday_label(p.deadline, today), source_names.get(p.source_id, p.source_id),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="center")
        title_cell = ws.cell(row=row, column=6)
        title_cell.hyperlink = p.url
        title_cell.font = LINK_FONT
        left = days_left(p.deadline, today)
        dday_cell = ws.cell(row=row, column=9)
        dday_cell.alignment = Alignment(horizontal="center")
        if left is not None and 0 <= left <= 1:
            dday_cell.fill = URGENT_FILL
        elif left is not None and 0 <= left <= 3:
            dday_cell.fill = SOON_FILL
    if postings:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(POSTING_COLUMNS))}{len(postings) + 1}"
    else:
        ws.cell(row=2, column=1, value="해당 공고가 없습니다.")


def _write_status(ws, results: list[SourceResult]) -> None:
    _header(ws, STATUS_COLUMNS)
    for n, r in enumerate(results, start=2):
        values = [r.name, r.org_type, r.state, r.fetched, r.matched, r.new, r.error, r.url]
        for col, value in enumerate(values, start=1):
            ws.cell(row=n, column=col, value=value)
        if r.is_problem:
            ws.cell(row=n, column=3).fill = PROBLEM_FILL
        elif r.is_unreachable:
            ws.cell(row=n, column=3).fill = UNREACHABLE_FILL
        if r.url:
            ws.cell(row=n, column=8).hyperlink = r.url
            ws.cell(row=n, column=8).font = LINK_FONT
