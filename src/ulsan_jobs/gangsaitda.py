"""강사잇다 '공고 올리기' 양식(엑셀)으로 내보내기.

config/gangsaitda_template.xlsx (강사잇다 양식 원본)을 그대로 열어 '공고' 시트에 마감 전 공고를 한 줄씩 채운다.
'예시'·'안내' 시트와 칸 이름·설명(메모)은 원본 그대로 둔다. 칸은 이름으로 찾으므로 양식의 칸 순서가 바뀌어도 된다.

필수 칸을 공고에서 찾지 못하면
    수업 일정   '원문 공고 참고'
    상세 내용   원문 제목·분야·접수 마감·원문 링크로 쓴 안내
    마감일      비우고 '처리' 칸에 '보류' (강사잇다가 그 줄을 올리지 않음. 확인해 채운 뒤 '보류'를 지우면 됨)
'메모' 칸에는 출처 게시판과 확인할 점을 적는다 (강사잇다는 메모 칸을 무시함).
"""
from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

from .classify import infer_district
from .models import Posting

TEMPLATE_NAME = "gangsaitda_template.xlsx"
SHEET = "공고"
HOLD = "보류"
SCHEDULE_FALLBACK = "원문 공고 참고"
_DISTRICTS = ("중구", "남구", "동구", "북구", "울주군")
_EXTRA_COLUMNS = {
    "처리": (10, "비어 있으면 올라가요. '보류'는 마감일을 찾지 못한 공고예요. 원문에서 확인해 마감일을 적고 '보류'를 지우세요."),
    "메모": (50, "강사잇다에는 올라가지 않는 참고 칸이에요 (수집한 게시판, 확인할 점)."),
}

_YEAR = re.compile(r"(?<!\d)(?:19|20)\d{2}\s*(?:학년도|년도|년|\.(?!\s*\d))\s*")  # '2026학년도', '2026. 방과후'
_PAREN_DATE = re.compile(r"\s*\([^()]*\d{1,2}\s*[./]\s*\d{1,2}[^()]*\)")
_TRAILING_RANGE = re.compile(r"\s+\d{1,2}\s*[./]\s*\d{1,2}\.?\s*[~\-–]\s*\d{1,2}\s*[./]\s*\d{1,2}\.?\s*$")


def clean_title(title: str) -> str:
    """양식 안내대로 제목에서 날짜·기간을 뺀다 ('2026학년도', '(~9.30까지)', '10.23-11.20')."""
    text = _TRAILING_RANGE.sub("", _PAREN_DATE.sub("", _YEAR.sub("", title or "")))
    text = re.sub(r"\]\s*", "] ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text if len(text) >= 5 else (title or "").strip()


def region(p: Posting) -> str:
    district = p.district if p.district in _DISTRICTS else infer_district(f"{p.org_name} {p.title}")
    return f"울산 {district}" if district else "울산"


def _source_label(source_name: str) -> str:
    return re.sub(r"\s+\([^()]*\)$", "", source_name or "")


def detail_text(p: Posting) -> str:
    """상세 내용: 원문 제목·분야·접수 마감·원문 링크로 쓴 안내."""
    lines = [p.title]
    if p.category and p.category != "기타":
        lines.append(f"분야: {p.category}")
    if p.deadline:
        lines.append(f"접수 마감: {p.deadline.isoformat()}")
    lines.append("자세한 내용과 지원 방법은 원문 공고를 확인해 주세요.")
    lines.append(f"원문 공고: {p.url}")
    return "\n".join(lines)


def row_values(p: Posting, source_name: str) -> dict[str, object]:
    info = p.info or {}
    schedule = info.get("schedule")
    notes = [f"출처: {_source_label(source_name)}"] if source_name else []
    if p.deadline is None:
        notes.insert(0, "마감일을 찾지 못해 보류: 원문에서 마감일을 확인해 적고 '보류'를 지우세요")
    if not schedule:
        notes.append("수업 일정은 원문 확인 필요")
    return {
        "제목": clean_title(p.title),
        "기관명": p.org_name or _source_label(source_name).split(" - ")[0],
        "지역": region(p),
        "마감일": p.deadline.isoformat() if p.deadline else "",
        "수업 일정": schedule or SCHEDULE_FALLBACK,
        "상세 내용": detail_text(p),
        "수업 대상": info.get("target", ""),
        "모집 인원": info.get("headcount"),
        "지원 자격": info.get("qualification", ""),
        "제출 서류": info.get("documents", ""),
        "원문 링크": p.url,
        "지원서 링크": "",
        "지원 이메일": info.get("email", ""),
        "처리": "" if p.deadline else HOLD,
        "메모": " · ".join(notes),
    }


def build_gangsaitda(
    path: Path, postings: list[Posting], source_names: dict[str, str], template: Path
) -> Path:
    """양식 원본을 복사해 공고를 채운 파일을 path 에 저장한다 (postings 순서대로)."""
    rows = [row_values(p, source_names.get(p.source_id, "")) for p in postings]
    return write_rows(path, rows, template)


def write_rows(path: Path, rows: list[dict[str, object]], template: Path) -> Path:
    """양식 원본을 복사해 rows(칸 이름 → 값)를 '공고' 시트에 한 줄씩 채워 path 에 저장한다."""
    wb = load_workbook(template)
    ws = wb[SHEET]
    columns = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    header_style = ws.cell(row=1, column=columns.get("수업 대상", ws.max_column))
    for name, (width, note) in _EXTRA_COLUMNS.items():
        if name in columns:
            continue
        col = ws.max_column + 1
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill, cell.alignment = copy(header_style.font), copy(header_style.fill), copy(header_style.alignment)
        cell.comment = Comment(note, "울산 강사공고 수집")
        ws.column_dimensions[cell.column_letter].width = width
        columns[name] = col

    body_font = Font(name=header_style.font.name, size=header_style.font.sz)
    wrap = Alignment(wrap_text=True, vertical="top")
    for n, values in enumerate(rows, start=2):
        for name, col in columns.items():
            value = values.get(name)
            cell = ws.cell(row=n, column=col, value=value if value not in ("", None) else None)
            cell.font, cell.alignment = body_font, wrap
            if name == "원문 링크" and value:
                cell.hyperlink = str(value)
    if rows:
        last = max(columns.values())
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=last).column_letter}{len(rows) + 1}"
    wb.active = wb.index(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_rows(data: bytes) -> list[dict[str, object]]:
    """다른 저장소가 만든 강사잇다 양식 파일의 '공고' 시트 → 줄마다 {칸 이름: 값} (빈 줄은 뺀다)."""
    from io import BytesIO

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(rows, ())]
    out = []
    for values in rows:
        row = {name: value for name, value in zip(header, values) if name and value not in (None, "")}
        if row.get("제목"):
            out.append(row)
    wb.close()
    return out
